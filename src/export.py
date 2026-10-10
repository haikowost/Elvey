"""v4 phase 5: exports (V4-SPEC.md §5).

The generic filtered CSV export (any column set, from whatever's currently shown in the People
Tree or Table view) is client-side JavaScript — `exportTableCsv()` in `src/static/index.html` —
since the browser already holds exactly the filtered/sorted rows the user is looking at.

This module is the other half: the **Quote template export**, a *named preset* of the same idea
that needs server-side support because its column set — the Consolidated Sales Contacts sheet's
own layout — pulls fields `/api/people` doesn't carry at all (ACCNO, Cluster, Axis Partner,
Milestone Partner, Installer Category, Sort). It's a round-trip: the same shape `src/ingest.py`
reads, written back out from the current DB state for whatever contacts are currently
filtered/selected in the dashboard, as an .xlsx (the sheet's native format — "where formatting
matters", per the spec).
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import shutil
import time
from pathlib import Path

import openpyxl

from . import db
from .config import load_config

# Exact column order of the Consolidated Sales Contacts sheet (src/ingest.py's FLAT_ALIASES).
QUOTE_COLUMNS = ("Ref", "Full Name", "Role", "Region/Branch", "Office Tel", "Cell number",
                 "Email Address", "Linkedin", "Company name", "ACCNO", "Axis Partner",
                 "Milestone Partner", "Installer Category", "AM", "Category", "In Zoho",
                 "In MakDB", "Cluster", "Sort")


def quote_rows(conn, contact_ids: list[int]) -> list[dict]:
    """One row per contact id, in QUOTE_COLUMNS shape, straight from the current DB state —
    including anything the harvester or a manual edit has since filled in. Unknown ids are
    skipped rather than erroring, since the id list comes from a client-side filtered view that
    could be stale by a request round-trip."""
    if not contact_ids:
        return []
    placeholders = ",".join("?" for _ in contact_ids)
    rows = db.rows(conn, f"""
        SELECT c.source_id, c.full_name, c.role, a.branch, c.tel, c.cell, c.email,
               COALESCE(c.linkedin_contact_url, c.linkedin_profile_url) AS linkedin,
               a.name AS company, a.accno, a.extra, c.allocated_rep, c.category, c.in_zoho,
               c.in_makdb, a.cluster, c.priority
        FROM contacts c LEFT JOIN accounts a ON a.id = c.account_id
        WHERE c.id IN ({placeholders})
        ORDER BY c.priority IS NULL, c.priority, c.full_name""", contact_ids)
    out = []
    for r in rows:
        extra = db.jload(r["extra"], {})
        ref = r["source_id"]
        if ref and ref.startswith("REF:"):
            ref = ref[4:]
        out.append({
            "Ref": ref, "Full Name": r["full_name"], "Role": r["role"], "Region/Branch": r["branch"],
            "Office Tel": r["tel"], "Cell number": r["cell"], "Email Address": r["email"],
            "Linkedin": r["linkedin"], "Company name": r["company"], "ACCNO": r["accno"],
            "Axis Partner": extra.get("axis_partner"), "Milestone Partner": extra.get("milestone_partner"),
            "Installer Category": extra.get("installer_category"), "AM": r["allocated_rep"],
            "Category": r["category"], "In Zoho": r["in_zoho"], "In MakDB": r["in_makdb"],
            "Cluster": r["cluster"], "Sort": r["priority"],
        })
    return out


def rows_to_xlsx(rows: list[dict], columns: tuple[str, ...] = QUOTE_COLUMNS, sheet_name: str = "Customers (import)") -> bytes:
    """Writes `rows` (dicts keyed by `columns`) to an .xlsx, one sheet, header row + data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(list(columns))
    for row in rows:
        ws.append([row.get(c) for c in columns])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- enriched round-trip

ENRICHED_EXTRA = ("LinkedIn Profile", "LinkedIn Summary", "Department", "Employment", "Current Company (LinkedIn)",
                  "City", "Province", "Country", "Photo File", "KYC Status", "Enriched At")


def enriched_rows(conn) -> list[dict]:
    """Every active contact in the Consolidated Sales Contacts column layout (so it re-imports
    cleanly), plus what the LinkedIn harvest added."""
    ids = [r["id"] for r in db.rows(conn, "SELECT id FROM contacts WHERE contact_status = 'active'")]
    base = quote_rows(conn, ids)
    extra = {r["full_name"] + "|" + (r["email"] or ""): r for r in db.rows(conn, """
        SELECT full_name, email, linkedin_profile_url, linkedin_summary, department, employment_status,
               linkedin_current_company, city, province, country, image_filename, image_status,
               enrich_status, enrich_last_at FROM contacts WHERE contact_status = 'active'""")}
    for row in base:
        e = extra.get(f"{row['Full Name']}|{row['Email Address'] or ''}") or {}
        has_photo = e.get("image_status") in ("downloaded", "manual")
        row.update({"LinkedIn Profile": e.get("linkedin_profile_url"), "LinkedIn Summary": e.get("linkedin_summary"),
                    "Department": e.get("department"), "Employment": e.get("employment_status"),
                    "Current Company (LinkedIn)": e.get("linkedin_current_company"), "City": e.get("city"),
                    "Province": e.get("province"), "Country": e.get("country"),
                    "Photo File": e.get("image_filename") if has_photo else None,
                    "KYC Status": e.get("enrich_status"), "Enriched At": e.get("enrich_last_at")})
        if row.get("Role") is None:
            row["Role"] = "lookup pending (KYC app)"  # keep the sheet's own placeholder on round-trip
    return base


def export_enriched(conn, cfg, out_dir: Path | None = None) -> dict:
    """data/exports/Elvey KYC enriched contacts <date>.xlsx + data/exports/faces/<face files>."""
    out_dir = Path(out_dir or cfg.data_dir / "exports")
    faces_out = out_dir / "faces"
    faces_out.mkdir(parents=True, exist_ok=True)
    rows = enriched_rows(conn)
    xlsx = out_dir / f"Elvey KYC enriched contacts {time.strftime('%Y%m%d')}.xlsx"
    xlsx.write_bytes(rows_to_xlsx(rows, QUOTE_COLUMNS + ENRICHED_EXTRA))
    copied = missing = 0
    src = cfg.kyc_folder
    for r in rows:
        name = r.get("Photo File")
        if not name:
            continue
        f = (src / name) if src else None
        if f and f.is_file():
            target = faces_out / name
            if not target.exists() or target.stat().st_mtime < f.stat().st_mtime:
                shutil.copy2(f, target)
            copied += 1
        else:
            missing += 1
    return {"xlsx": str(xlsx), "contacts": len(rows), "faces_dir": str(faces_out), "faces": copied,
            "faces_missing": missing}


# --------------------------------------------------------------------------- people snapshot

SNAPSHOT_JSON = "people_snapshot.json"
SNAPSHOT_CSV = "people_snapshot.csv"
SNAPSHOT_COLUMNS = ("id", "name", "company", "segment", "tier", "score", "kyc_order", "pinned", "role", "department",
                    "email", "linkedin", "location", "employment", "now_at", "kyc_status", "complete",
                    "last_harvested", "correction", "summary")
THUMB_PX = 72


def snapshot_dir(cfg) -> Path:
    """Where people_snapshot.json goes: exports.snapshot_dir when its folder (or its parent) exists;
    blank = the Google Drive KYC folder (paths.kyc_folder, a Drive for desktop path) when that exists;
    else data/exports. Relative paths resolve against config.yaml."""
    configured = cfg.path("exports", "snapshot_dir")
    if configured and (configured.exists() or configured.parent.exists()):
        return configured
    if not configured and cfg.kyc_folder and cfg.kyc_folder.is_dir():
        return cfg.kyc_folder
    return cfg.data_dir / "exports"


def _thumb(cfg, filename: str | None, px: int = THUMB_PX) -> str | None:
    if not filename or not cfg.kyc_folder:
        return None
    f = cfg.kyc_folder / filename
    if not f.is_file():
        return None
    try:
        from PIL import Image

        img = Image.open(f).convert("RGB")
        img.thumbnail((px, px))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=70, optimize=True)
        return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
    except Exception:  # noqa: BLE001 - one unreadable face must not stop the snapshot
        return None


def snapshot_data(conn, cfg) -> dict:
    """Everything the people view needs, as plain JSON: people (tier, score, role, company, profile URL,
    small photo, last harvested, correction notes), accounts and the links between them."""
    from .harvest import REFRESH_DAYS, is_complete
    from .priority import tier_counts

    days = int((cfg.get("harvest") or {}).get("refresh_days", REFRESH_DAYS))
    roles: dict[int, list[str]] = {}
    for r in db.rows(conn, "SELECT account_id, role FROM account_roles ORDER BY role"):
        roles.setdefault(r["account_id"], []).append(r["role"])
    accounts = [{"id": f"a{a['id']}", "name": a["name"], "segment": a["segment"], "rank": a["rank"], "rep": a["rep"],
                 "division": a["division"], "kyc_status": a["kyc_status"], "roles": roles.get(a["id"], [])}
                for a in db.rows(conn, "SELECT * FROM accounts ORDER BY rank IS NULL, rank, name")]
    people, links = [], []
    for c in db.rows(conn, """SELECT c.*, a.name AS company FROM contacts c LEFT JOIN accounts a ON a.id = c.account_id
                              WHERE c.contact_status = 'active'
                              ORDER BY c.kyc_priority_order IS NULL, c.kyc_priority_order, c.priority IS NULL,
                                       c.priority, c.full_name"""):
        has_photo = c["image_status"] in ("downloaded", "manual")
        people.append({
            "id": f"c{c['id']}", "name": c["full_name"], "company": c["company"],
            "account_id": f"a{c['account_id']}" if c["account_id"] else None, "segment": c["segment"],
            "tier": c["kyc_tier"], "score": c["kyc_score"], "kyc_order": c["kyc_priority_order"],
            "pinned": bool(c["kyc_pinned"]), "role": c["role"], "department": c["department"], "email": c["email"],
            "linkedin": c["linkedin_profile_url"] or (c["linkedin_contact_url"] if c["linkedin_contact_url"] and
                                                      "/in/" in c["linkedin_contact_url"] else None),
            "location": ", ".join(x for x in (c["city"], c["province"], c["country"]) if x) or None,
            "employment": c["employment_status"], "now_at": c["linkedin_current_company"],
            "kyc_status": c["enrich_status"], "complete": is_complete(c, days),
            "last_harvested": c["enrich_last_at"], "correction": c["kyc_correction"],
            "summary": c["linkedin_summary"], "experience": db.jload(c["linkedin_experience"], []),
            "photo": _thumb(cfg, c["image_filename"]) if has_photo else None,
        })
        if c["account_id"]:
            links.append({"source": f"c{c['id']}", "target": f"a{c['account_id']}", "rel": "works_at"})
        if c["reports_to_id"]:
            links.append({"source": f"c{c['id']}", "target": f"c{c['reports_to_id']}", "rel": "reports_to"})
    for l in db.rows(conn, "SELECT source, target, rel FROM graph_links"):
        links.append(dict(l))
    last_run = db.jload((db.one(conn, "SELECT value FROM meta WHERE key = 'last_harvest_run'") or {}).get("value"))
    return {"generated_at": time.strftime("%Y-%m-%dT%H:%M:%S"), "source": "elvey-kyc local app",
            "counts": {"people": len(people), "accounts": len(accounts), "links": len(links),
                       "with_photo": sum(1 for p in people if p["photo"]),
                       "complete": sum(1 for p in people if p["complete"])},
            "tiers": tier_counts(conn), "last_harvest_run": last_run,
            "people": people, "accounts": accounts, "links": links}


def write_snapshot(conn, cfg, out_dir: Path | str | None = None) -> dict:
    """people_snapshot.json + people_snapshot.csv into exports.snapshot_dir (written to a temp name
    and renamed, so Drive never syncs half a file)."""
    import csv

    out = Path(out_dir) if out_dir else snapshot_dir(cfg)
    out.mkdir(parents=True, exist_ok=True)
    data = snapshot_data(conn, cfg)
    js, cs = out / SNAPSHOT_JSON, out / SNAPSHOT_CSV
    tmp = js.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":"), default=str), encoding="utf-8")
    tmp.replace(js)
    tmp = cs.with_suffix(".csv.tmp")
    with open(tmp, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(SNAPSHOT_COLUMNS), extrasaction="ignore")
        w.writeheader()
        for p in data["people"]:
            w.writerow({**p, "pinned": "Y" if p["pinned"] else "", "complete": "Y" if p["complete"] else ""})
    tmp.replace(cs)
    with conn:
        conn.execute("INSERT OR REPLACE INTO meta(key, value) VALUES ('last_snapshot', ?)",
                     [json.dumps({"at": data["generated_at"], "json": str(js), "people": len(data["people"])})])
    return {"json": str(js), "csv": str(cs), **data["counts"]}


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description="Export the enriched contacts (.xlsx) + face images, or "
                                             "`snapshot`: people_snapshot.json/.csv for the people view")
    ap.add_argument("cmd", nargs="?", choices=("enriched", "snapshot"), default="enriched")
    ap.add_argument("--out", help="output folder (default data/exports; snapshot: exports.snapshot_dir)")
    ap.add_argument("--config")
    args = ap.parse_args(argv)
    cfg = load_config(args.config)
    conn = db.connect(cfg.db_path)
    if args.cmd == "snapshot":
        print(json.dumps(write_snapshot(conn, cfg, args.out), indent=2))
        return
    print(json.dumps(export_enriched(conn, cfg, args.out), indent=2))


if __name__ == "__main__":
    main()
