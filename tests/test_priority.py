"""KYC contact priority: import mapping (CSV + reviewed workbook), idempotency, harvest order."""
from __future__ import annotations

import csv

import openpyxl
from fastapi.testclient import TestClient

from src import dashboard, db, graph, harvest, ingest, priority

HEADER = ["kyc_order", "contact_id", "account_id", "full_name", "company", "email", "tier", "score"]
ROWS = [
    [1, "C1", "A0001", "Anna Smith", "Acme Security (Pty) Ltd", "ANNA@acme.co.za", "P1 Key", 80.4],  # email (case)
    [2, "C4", "A0002", "Carla Müller", "Beta Integrators", "", "P1 Key", 70],                        # contact_id
    [3, "ZZZ", "A0003", "Dan Out", "Gamma Systems", "", "P2 Active", 50],                            # name + company
    [4, "C2", "A0001", "Bob Jones", "Acme Security (Pty) Ltd", "bob@acme.co.za", "P4 Low relevance", 20],
    [5, "C8", "", "Leandro da Cunha", "Duxbury", "", "Internal/Competitor (curated)", 0],
    [6, "C99", "A0009", "Nobody Known", "Nowhere", "x@nowhere.co.za", "P3 Reference", 40],
]


def _csv(path):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(HEADER)
        w.writerows(ROWS)
    return path


def _workbook(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("All contacts scored")
    ws.append(["KYC order", "Tier", "Score (0–100)", "Name", "Company", "Email", "contact_id"])
    for r in ROWS:
        ws.append([r[0], r[6], r[7], r[3], r[4], r[5] or None, r[1]])
    ws = wb.create_sheet("P1 Key — confirm")
    ws.append(["KYC order", "Score (0–100)", "Name", "Company", "Email", "contact_id",
               "Confirm (Y / N / Not relevant)", "Correction (name, role, employer, email…)"])
    ws.append([1, 80.4, "Anna Smith", "Acme Security (Pty) Ltd", "anna@acme.co.za", "C1", "Not relevant", None])
    ws.append([2, 70, "Carla Müller", "Beta Integrators", None, "C4", "y", None])
    ws = wb.create_sheet("P2 Active")
    ws.append(["KYC order", "Score (0–100)", "Name", "Company", "Email", "contact_id",
               "Confirm (Y / N / Not relevant)", "Correction (name, role, employer, email…)"])
    ws.append([3, 50, "Dan Out", "Gamma Systems", None, "ZZZ", "N", "Now at Beta Integrators"])
    wb.save(path)
    return path


def _tiers(conn):
    return {r["full_name"]: r for r in db.rows(conn, "SELECT * FROM contacts")}


def test_csv_import_matching_and_mapping(cfg, conn, tmp_path):
    ingest.run(cfg, all_rows=True)
    assert not priority.has_priority(conn)
    s = priority.import_file(conn, _csv(tmp_path / "p.csv"))
    assert s["scored"] == 6 and s["matched"] == 5 and s["unmatched_names"] == ["Nobody Known"]
    t = _tiers(conn)
    assert (t["Anna Smith"]["kyc_tier"], t["Anna Smith"]["kyc_score"], t["Anna Smith"]["kyc_priority_order"]) == ("P1 Key", 80.4, 1)
    assert t["Carla Müller"]["kyc_tier"] == "P1 Key" and t["Dan Out"]["kyc_tier"] == "P2 Active"
    assert t["Leandro da Cunha"]["kyc_tier"] == "Internal/Competitor (curated)"
    assert t["Gordon Moore"]["kyc_tier"] is None
    again = priority.import_file(conn, tmp_path / "p.csv")
    assert again["matched"] == 5 and _tiers(conn)["Anna Smith"]["kyc_tier"] == "P1 Key"


def test_reviewed_workbook_y_n_not_relevant_and_corrections(cfg, conn, tmp_path):
    ingest.run(cfg, all_rows=True)
    s = priority.import_file(conn, _workbook(tmp_path / "Elvey KYC Contact Priority 20261010.xlsx"))
    assert (s["pinned"], s["demoted"], s["excluded"], s["corrections"]) == (1, 1, 1, 1)
    t = _tiers(conn)
    assert t["Anna Smith"]["kyc_tier"] == "Excluded" and t["Anna Smith"]["kyc_tier_base"] == "P1 Key"
    assert t["Carla Müller"]["kyc_pinned"] == 1 and t["Carla Müller"]["kyc_tier"] == "P1 Key"
    dan = t["Dan Out"]
    assert dan["kyc_tier"] == "P3 Reference" and dan["kyc_correction"] == "Now at Beta Integrators"
    assert dan["account_id"] == db.one(conn, "SELECT id FROM accounts WHERE name = 'Gamma Systems'")["id"]  # not auto-moved
    # re-importing the plain CSV afterwards keeps the review's effect (idempotent in any order)
    priority.import_file(conn, _csv(tmp_path / "p.csv"))
    t = _tiers(conn)
    assert t["Anna Smith"]["kyc_tier"] == "Excluded" and t["Dan Out"]["kyc_tier"] == "P3 Reference"


def test_normalise_review():
    assert priority.normalise_review(" Yes ") == ("Y", None)
    assert priority.normalise_review("n") == ("N", None)
    assert priority.normalise_review("Not relevant") == ("Not relevant", None)
    assert priority.normalise_review("wrong person - it's Jan") == (None, "wrong person - it's Jan")
    assert priority.normalise_review(None) == (None, None)


def test_harvest_order_with_priority(cfg, conn, tmp_path):
    ingest.run(cfg, all_rows=True)
    legacy = [r["full_name"] for r in harvest.queue(conn, cfg)]
    assert legacy[0] == "Anna Smith"  # no priority data: account-rank order still works
    priority.import_file(conn, _csv(tmp_path / "p.csv"))
    names = [r["full_name"] for r in harvest.queue(conn, cfg)]
    # curated never-harvested first, then P1 by order, then P2, then unscored; P4 (Bob) skipped
    assert names[:4] == ["Leandro da Cunha", "Anna Smith", "Carla Müller", "Dan Out"]
    assert "Bob Jones" not in names and set(names[4:]) == {"Jaco Moolman", "Gordon Moore"}
    assert [r["full_name"] for r in harvest.queue(conn, cfg, tier="p1")] == names[:3]
    assert [r["full_name"] for r in harvest.queue(conn, cfg, tier="top500")] == names[:4]  # legacy name = p2
    assert harvest.queue(conn, cfg, include_low=True)[-1]["full_name"] == "Bob Jones"
    bob = _tiers(conn)["Bob Jones"]["id"]
    assert [r["id"] for r in harvest.queue(conn, cfg, ids=[bob])] == [bob]  # explicit ids always honoured

    # Y pins Carla ahead of the rest of P1; a curated contact already harvested drops behind P3
    priority.import_file(conn, _workbook(tmp_path / "rev.xlsx"))
    conn.execute("UPDATE contacts SET enrich_last_at = datetime('now'), enrich_attempts = 1, enrich_status = 'failed' "
                 "WHERE full_name = 'Leandro da Cunha'")
    conn.commit()
    names = [r["full_name"] for r in harvest.queue(conn, cfg)]
    assert names[0] == "Carla Müller" and names[1] == "Dan Out" and "Anna Smith" not in names
    assert "Leandro da Cunha" in names[2:]


def test_progress_people_tree_and_graph_expose_tiers(cfg, conn, tmp_path):
    ingest.run(cfg, all_rows=True)
    priority.import_file(conn, _csv(tmp_path / "p.csv"))
    c = TestClient(dashboard.create_app(cfg))
    p = c.get("/api/kyc/progress").json()
    assert p["priority_mode"] is True
    tiers = {t["tier"]: t for t in p["tiers"]}
    assert p["tiers"][0]["tier"] == "Internal/Competitor (curated)"
    assert tiers["P1 Key"]["total"] == 2 and tiers["P1 Key"]["pending"] == 2 and tiers["P1 Key"]["enriched"] == 0
    tree = c.get("/api/people").json()
    acme = next(g for g in tree["customer"] if g["title"] == "Acme Security (Pty) Ltd")
    assert [x["name"] for x in acme["people"]][0] == "Anna Smith"  # highest score first
    assert acme["people"][0]["kyc_tier"] == "P1 Key" and acme["people"][0]["kyc_score"] == 80.4
    g = graph.load_live(conn, cfg)
    anna = next(n for n in g.nodes if n["label"] == "Anna Smith")
    assert anna["kycTier"] == "P1 Key" and anna["kycScore"] == 80.4 and "KYC P1 Key · score 80.4" in anna["facts"]
    acme_id = next(n["id"] for n in g.nodes if n["label"] == "Acme Security (Pty) Ltd")
    assert graph.investigate(g, acme_id)["contacts"][0]["label"] == "Anna Smith"


def test_default_inputs_and_cli(cfg, conn, tmp_path, capsys):
    ingest.run(cfg, all_rows=True)
    cfg["inputs"]["kyc_priority"] = str(_workbook(tmp_path / "Elvey KYC Contact Priority 20261010.xlsx"))
    files = priority.default_inputs(cfg)
    assert files[0] == priority.SEED_CSV and files[-1].name.endswith("20261010.xlsx")
