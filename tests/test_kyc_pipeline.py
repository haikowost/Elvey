"""The relational-workbook pipeline end to end: facts ingest (sellout / quotes / open Zoho deals),
the live graph built from them, Spider defaults (root + trimmed ego), the prioritised harvest queue,
the name+company search fallback, the KYC progress endpoint and the enriched export."""
from __future__ import annotations

import datetime as dt
import os

import openpyxl
from fastapi.testclient import TestClient

from src import dashboard, db, export, graph, harvest, images, ingest
from src.harvest import LinkedInDriver, Result

from .conftest import SHEETS
from .test_harvest import FakeDriver, good

FACTS = {
    "fact_quotes": [["quote_id", "q_no", "rep", "date", "for_contact", "ref", "account_id", "value", "gp"],
                    [1, "SN100", "Siresen", dt.datetime(2026, 9, 1), "Anna", "Tower CCTV", "A0001", 125000, None],
                    [2, "SN101", "Siresen", dt.datetime(2026, 9, 20), "Bob", "Tower phase 2", "A0001", None, None],
                    [3, "HW200", "Haiko", dt.datetime(2026, 8, 2), "Carla", "Beta MS", "A0002", 5000, None],
                    [4, "HW201", "Haiko", dt.datetime(2026, 8, 3), "?", "Walk-in", None, 100, None]],
    "fact_deals": [["deal_id", "deal_name", "account_id", "amount", "closing_date", "stage", "open_flag"],
                   [1, "Acme HQ Milestone upgrade", "A0001", 72609.28, dt.datetime(2026, 10, 30), None, None],
                   [2, "Acme old job", "A0001", 1852, dt.datetime(2026, 1, 30), "Closed Won", "open"],
                   [3, "Lost cause", "A0002", 99, None, "Closed Lost to Competition", "open"],
                   [4, "Orphan deal", "A9999", 10, None, None, None]],
}


def write_relational(path, extra=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in {**SHEETS, **FACTS, **(extra or {})}.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def _relational_cfg(cfg, tmp_path):
    cfg["inputs"]["workbook"] = str(write_relational(tmp_path / "Elvey Consolidated Relational Database 20260930.xlsx"))
    return cfg


def test_relational_facts_ingest_idempotent(cfg, conn, tmp_path):
    _relational_cfg(cfg, tmp_path)
    out = ingest.run(cfg)
    assert out["format"] == "relational"
    f = out["facts"]
    assert f["quotes"] == 4 and f["open_deals"] == 1 and f["deals_unmatched"] == 1 and f["sellout"] == 2
    acme = db.one(conn, "SELECT id FROM accounts WHERE name = 'Acme Security (Pty) Ltd'")["id"]
    assert {r["brand"]: r["fy26"] for r in db.rows(conn, "SELECT * FROM sellout WHERE account_id = ?", [acme])} == \
        {"Bosch": 100, "Milestone": 900}
    deal = db.one(conn, "SELECT * FROM opportunities WHERE source = 'zoho_deal'")
    assert deal["title"] == "Acme HQ Milestone upgrade" and deal["value"] == "R 72,609" and deal["account_id"] == acme
    assert deal["closing_date"] == "2026-10-30" and deal["stage"] == "Open"
    assert db.one(conn, "SELECT count(*) n FROM quotes WHERE account_id = ?", [acme])["n"] == 2
    assert db.one(conn, "SELECT count(*) n FROM quotes WHERE account_id IS NULL")["n"] == 1  # kept, just unlinked
    # contacts carry their account's resolved owner + call cadence, and the top-500 rank for the harvester
    anna = db.one(conn, "SELECT * FROM contacts WHERE full_name = 'Anna Smith'")
    assert anna["allocated_rep"] == "SN" and anna["category"] == "A" and anna["top500_rank"] == 1

    again = ingest.run(cfg)
    assert set(again["upsert"]) <= {"accounts_unchanged", "contacts_unchanged"}
    for t, n in (("quotes", 4), ("sellout", 2), ("opportunities", 1)):
        assert db.one(conn, f"SELECT count(*) n FROM {t}")["n"] == n

    # the deal closes in Zoho -> next ingest drops it from the open opportunities
    closed = [r[:] for r in FACTS["fact_deals"]]
    closed[1][5] = "Closed Won"
    write_relational(cfg.path("inputs", "workbook"), {"fact_deals": closed})
    assert ingest.run(cfg)["facts"]["deals_closed_removed"] == 1
    assert db.one(conn, "SELECT count(*) n FROM opportunities")["n"] == 0


def test_relational_then_flat_keeps_ranks_and_elvey_segment(cfg, conn, tmp_path):
    """Both configured inputs, relational first: the flat sheet must not wipe the top-50 rank or
    re-label Elvey (no Segment column) as a customer."""
    rel = write_relational(tmp_path / "rel.xlsx")
    flat = tmp_path / "flat.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers (import)"
    ws.append(["Ref", "Full Name", "Role", "Email Address", "Linkedin", "Company name", "AM", "Category"])
    ws.append(["AnnaS", "Anna Smith", "lookup pending (KYC app)", "anna@acme.co.za", "in ›", "Acme Security", "SN", "A"])
    ws.append(["TashaS", "Tasha Smith", "Sales", "tasha@elvey.co.za", "in ›", "Elvey", "1house", "A"])
    wb.save(flat)
    cfg["inputs"].update({"relational_workbook": str(rel), "workbook": str(flat), "sheet": "Customers (import)"})
    out = ingest.run(cfg)
    assert [r["format"] for r in out["runs"]] == ["relational", "flat"]
    acme = db.one(conn, "SELECT * FROM accounts WHERE name_norm = 'acme security'")
    assert acme["rank"] == 1
    assert db.one(conn, "SELECT segment FROM accounts WHERE name = 'Elvey'")["segment"] == "internal"
    assert db.one(conn, "SELECT segment FROM contacts WHERE full_name = 'Tasha Smith'")["segment"] == "internal"
    anna = db.one(conn, "SELECT * FROM contacts WHERE full_name = 'Anna Smith'")
    # 'in ›' is not a URL, and a 'lookup pending' placeholder never overwrites a real role from the workbook
    assert anna["linkedin_contact_url"] is None and anna["role"] == "CEO"


def test_latest_like_picks_newest_dated_workbook(tmp_path):
    old = tmp_path / "Elvey Consolidated Relational Database 20260924.xlsx"
    new = tmp_path / "Elvey Consolidated Relational Database 20260930.xlsx"
    for p, t in ((old, 1_000_000), (new, 2_000_000)):
        p.write_bytes(b"x")
        os.utime(p, (t, t))
    (tmp_path / "~$Elvey Consolidated Relational Database 20261001.xlsx").write_bytes(b"lock")
    assert ingest.latest_like(tmp_path / "Elvey Consolidated Relational Database 20261015.xlsx") == new
    assert ingest.latest_like(tmp_path / "Elvey Consolidated Relational Database.xlsx") == new
    assert ingest.latest_like(old) == old  # an existing file is used as configured


def test_live_graph_from_relational_data(cfg, conn, tmp_path):
    _relational_cfg(cfg, tmp_path)
    ingest.run(cfg)
    g = graph.load_live(conn, cfg)
    elvey = next(n for n in g.nodes if n["label"] == "Elvey")
    assert g.root == elvey["id"] and g.payload()["root"] == elvey["id"]
    acme = next(n for n in g.nodes if n["label"] == "Acme Security (Pty) Ltd")
    # rep (account owner) by name, quotes as interaction weight
    link = next(l for l in g.links if l["source"] == elvey["id"] and l["target"] == acme["id"])
    assert link["rel"] == "does_business_with"
    assert link["ex"] == {"owner": "Siresen Naidoo", "quotes": 2, "last_quote": "2026-09-20"}
    assert "Quotes: 2 (last 2026-09-20)" in acme["facts"]
    # open Zoho deal -> opportunity; brands from fact_sellout -> product nodes
    assert acme["oppCount"] == 1 and acme["opps"][0]["t"] == "Acme HQ Milestone upgrade"
    milestone = g.by_id["brand:milestone"]
    assert milestone["type"] == "product" and milestone["sub"] == "Portfolio brand"
    assert any(l["source"] == acme["id"] and l["target"] == "brand:milestone" and l["ex"]["fy26"] == 900 for l in g.links)
    assert any(l["source"] == "brand:milestone" and l["target"] == elvey["id"] and l["rel"] == "part_of" for l in g.links)


def test_ego_trims_big_hubs_and_reports_totals(cfg, conn, tmp_path):
    _relational_cfg(cfg, tmp_path)
    ingest.run(cfg, all_rows=True)
    g = graph.load_live(conn, cfg)
    full = graph.ego(g, g.root, 1, limit=0)
    dbw = next(gr for gr in full["groups"] if gr["rel"] == "does_business_with")
    assert dbw["total"] == len(dbw["ids"]) >= 2
    top = graph.ego(g, g.root, 1, limit=1)
    dbw1 = next(gr for gr in top["groups"] if gr["rel"] == "does_business_with")
    assert len(dbw1["ids"]) == 1 and dbw1["total"] == dbw["total"]
    # the biggest / most-quoted / open-deal account wins the one slot
    assert g.by_id[dbw1["ids"][0]]["label"] == "Acme Security (Pty) Ltd"
    assert {n["id"] for n in top["nodes"]} == {i for gr in top["groups"] for i in gr["ids"]}


def test_demo_seed_only_behind_env_flag(cfg, monkeypatch):
    monkeypatch.delenv("ELVEY_GRAPH_SOURCE", raising=False)
    cfg["graph"] = {"source": "seed"}
    assert graph.source_name(cfg) == "live"  # a config file alone can't switch real data to demo data
    monkeypatch.setenv("ELVEY_GRAPH_SOURCE", "seed")
    assert graph.source_name(cfg) == "seed"
    assert graph.load_seed().root == "elvey"


def test_harvest_queue_priority_tiers(cfg, conn, tmp_path):
    _relational_cfg(cfg, tmp_path)
    cfg["scope"].update({"top_accounts": None, "top_contacts": None})
    ingest.run(cfg)
    q = harvest.queue(conn, cfg)
    tiers = [r["tier"] for r in q]
    assert tiers == sorted(tiers)
    names = [r["full_name"] for r in q]
    # Acme (rank 1) people first, then Beta (rank 2), Gamma (rank 3), then Elvey & competitors
    assert names[:2] == ["Anna Smith", "Bob Jones"] and names[2:4] == ["Carla Müller", "Dan Out"]
    assert q[0]["tier"] == 0 and next(r for r in q if r["full_name"] == "Leandro da Cunha")["tier"] == 3
    top50 = {r["full_name"] for r in harvest.queue(conn, cfg, tier="top50")}
    assert top50 == {"Anna Smith", "Bob Jones", "Carla Müller", "Dan Out"}
    assert len(harvest.queue(conn, cfg, tier="all")) == len(q)


def test_is_profile_url():
    assert harvest.is_profile_url("https://www.linkedin.com/in/bobjones")
    assert harvest.is_profile_url("https://za.linkedin.com/in/bob-jones-1a2b3c/?trk=x")
    for bad in (None, "", "in ›", "https://www.linkedin.com/company/acme", "https://www.linkedin.com/in/",
                "https://www.linkedin.com/search/results/people/?keywords=bob"):
        assert not harvest.is_profile_url(bad), bad


class _Page:
    def __init__(self):
        self.url = ""

    def evaluate(self, *a):
        return None


def test_fetch_searches_name_and_company_when_url_is_a_placeholder():
    d = object.__new__(LinkedInDriver)
    d.h, d.page = {}, _Page()
    searched = []
    d._search = lambda c: (searched.append(c["full_name"]), (None, "no result", "dbg"))[1]
    d._wait = lambda *a, **k: None
    for stored in (None, "in ›", "https://www.linkedin.com/company/acme"):
        res = d.fetch({"full_name": "Anna Smith", "company": "Acme", "linkedin_contact_url": stored}, need_face=False)
        assert res.status == "no_profile"
    assert searched == ["Anna Smith"] * 3


def test_run_records_last_harvest_and_progress_endpoint(cfg, loaded):
    before = TestClient(dashboard.create_app(cfg)).get("/api/kyc/progress").json()
    assert before["last_harvest_run"] is None and before["linkedin_url"] == 1  # Bob's real URL from the sheet
    harvest.run(loaded, cfg, lambda: FakeDriver({"Anna Smith": good("Anna Smith")}), limit=2,
                sleep=lambda s: None, log=lambda *a: None, tier="all")
    p = TestClient(dashboard.create_app(cfg)).get("/api/kyc/progress").json()
    assert p["last_harvest_run"]["processed"] == 2 and p["last_harvest_run"]["tier"] == "all"
    assert p["used_today"] == 2 and p["daily_cap"] == 40 and p["enriched"] >= 1
    assert p["contacts"] == before["contacts"] and p["role_known"] >= before["role_known"]
    assert [t["tier"] for t in p["tiers"]] == ["top50", "top500", "cat-a"]
    assert p["last_linkedin_visit"]


def test_export_enriched_xlsx_and_faces(cfg, loaded, tmp_path):
    out = export.export_enriched(loaded, cfg, tmp_path / "exp")
    wb = openpyxl.load_workbook(out["xlsx"], read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    header = rows[0]
    assert header[:len(export.QUOTE_COLUMNS)] == export.QUOTE_COLUMNS and "LinkedIn Profile" in header
    assert len(rows) - 1 == out["contacts"] == db.one(loaded, "SELECT count(*) n FROM contacts")["n"]
    faces = sorted(p.name for p in (tmp_path / "exp" / "faces").iterdir())
    assert out["faces"] == len(faces) >= 1 and "CUST__Anna_Smith.jpg" in faces
    # the round-trip keeps the sheet's own placeholder for a role nobody has looked up yet
    role_col = header.index("Role")
    assert any(r[role_col] == "lookup pending (KYC app)" for r in rows[1:])


def test_images_match_still_runs_after_relational_ingest(cfg, conn, tmp_path):
    _relational_cfg(cfg, tmp_path)
    ingest.run(cfg)
    images.match(conn, cfg)
    assert db.one(conn, "SELECT image_status FROM contacts WHERE full_name = 'Anna Smith'")["image_status"] == "downloaded"


def test_browser_launch_failure_stops_cleanly(cfg, loaded):
    def boom():
        raise RuntimeError("Chromium distribution 'chrome' is not found\nmore detail")
    stats = harvest.run(loaded, cfg, boom, sleep=lambda s: None, log=lambda *a: None)
    assert stats["processed"] == 0 and stats["stopped"].startswith("browser could not start: Chromium")
